import openpyxl
import pandas as pd
import streamlit as st
from datetime import datetime, time, timedelta
import os
import urllib.parse
import base64

st.set_page_config(page_title="Wyszukiwarka Zastępstw", page_icon="📋", layout="wide")
st.title("Wyszukiwarka Zastępstw 📋")

# --- UKRYWANIE ANGIELSKICH KOMUNIKATÓW STREAMLIT ---
st.markdown("""
    <style>
    /* Ukrywa domyślny tekst 'Press Enter to apply' w polach tekstowych */
    div[data-testid="InputInstructions"] {
        display: none;
    }
    </style>
""", unsafe_allow_html=True)

# --- WYŚWIETLANIE ZASAD (PDF + TEKST) ---
with st.expander("📖 Przypomnij zasady znajdywania zastępstw"):
    pdf_file = "ZASTĘPSTWA 202627.pdf"
    if os.path.exists(pdf_file):
        with open(pdf_file, "rb") as f:
            pdf_bytes = f.read()
        st.download_button(label="📥 Pobierz oryginalny schemat (PDF)", data=pdf_bytes, file_name=pdf_file, mime="application/pdf")
    else:
        st.info(f"💡 Aby umożliwić pobieranie oryginalnego pliku, upewnij się, że plik `{pdf_file}` znajduje się w tym samym folderze co aplikacja.")

    # Wersja tekstowa zintegrowana w aplikacji
    st.markdown("""
    ### 🚨 Co robić, kiedy potrzebuję zastępstwa?
    
    #### 1️⃣ Sytuacja z "ZAPASEM CZASOWYM" (wiem z wyprzedzeniem):
    * Sprawdzam w pliku *ZASTĘPSTWA KUM&CO 2026/27*, kto nie uczy w czasie moich zajęć. (Albo odpalam sobie ten program, bo taką jestem spryciulą.)
    * W grupach **przedszkolnych i 0-3** szukam zastępstwa stacjonarnie.
    * W grupach **4+** mogę poszukać zastępstw na zajęcia **ON-LINE**, jeśli nie znajdę stacjonarnie *(min. 1 dzień wcześniej biuro musi wysłać informację do rodzica)*.
    * Wysyłam e-mail/SMS lub dzwonię **TYLKO DO OSÓB**, które nie uczą w trakcie moich zajęć. 
    * (Jeśli wyślesz maila do wszystkich **POMIMO** użyciu tego programu, to rozważę rozwiązania siłowe.) 
    
    **🟢 Znajduję zastępstwo ;)**
    W tabeli z zastępstwami (w filii, w której pracuję) wpisuję informacje o tym, co trzeba zrealizować (tylko zakres materiału), załączając w DW lidera, biuro i metodyków swojej filii. Biuro udostępnia e-dziennik na zastępstwo.
    
    *Obowiązki zastępcy:* Lektor prowadzący zastępstwo przygotowuje, prowadzi i podsumowuje całą lekcję (sprawdzając także short tests i tests, ale bez writings!), uzupełniając e-dziennik (oceny, obecności, ew. praca domowa) i informuje o niepokojących sytuacjach lektora głównego.

    ---
    #### 2️⃣ Sytuacja NAGŁA, LOSOWA (lub brak chętnych):
    Jeśli nie jesteś w stanie szukać zastępstwa LUB szukałeś/aś, ale **nie znajdujesz zastępstwa :(**
    * Kontaktujesz się z sekretariatem szkoły i liderem, informując o trudnościach.
    * Podajesz zakres materiału do zrealizowania na zajęciach.
    * Jeśli nie uda się przeprowadzić zajęć (lider i sekretariat też nie znajdą zastępstwa) – sekretariat odwołuje zajęcia, a Ty podajesz terminy na ich odrobienie.
    """)

# --- ZARZĄDZANIE PAMIĘCIĄ SESJI I CALLBACKI ---
if "search_results" not in st.session_state:
    st.session_state.search_results = None

# Pobranie imienia z pamięci URL (Ciasteczka URL)
if "user_name" not in st.session_state:
    st.session_state.user_name = st.query_params.get("name", "")

def update_name():
    st.query_params["name"] = st.session_state.user_name

def update_end_time(group_idx):
    start_key = f"start_{group_idx}"
    end_key = f"end_{group_idx}"
    if start_key in st.session_state:
        start_val = st.session_state[start_key]
        start_dt = datetime.combine(datetime.today(), start_val)
        st.session_state[end_key] = (start_dt + timedelta(minutes=90)).time()

def update_branches():
    if "branch_0" in st.session_state:
        new_branch = st.session_state["branch_0"]
        for i in range(1, 6):
            if f"branch_{i}" in st.session_state:
                st.session_state[f"branch_{i}"] = new_branch

def next_tpl(idx):
    st.session_state[f"tpl_{idx}"] = (st.session_state.get(f"tpl_{idx}", 0) + 1) % 6

# --- FUNKCJE I SŁOWNIKI BAZOWE ---
def is_cell_free(cell):
    if not cell.fill or cell.fill.fill_type is None: return True
    color = getattr(cell.fill.start_color, "index", None)
    if color in [0, "00000000", "FFFFFFFF", "None", None]: return True
    rgb = getattr(cell.fill.start_color, "rgb", None)
    if rgb in ["00000000", "FFFFFFFF", None]: return True
    return False

days = ["Poniedziałek", "Wtorek", "Środa", "Czwartek", "Piątek"]
levels = ["3-5 LAT", "0", "1", "2", "3", "4", "5", "6", "7", "8+", "MASTER"]
day_mapping = {"Poniedziałek": "PON", "Wtorek": "WT", "Środa": "ŚR", "Czwartek": "CZW", "Piątek": "PT"}

grammar_days = {
    "Poniedziałek": {"na": "poniedziałek", "w": "w poniedziałek", "najblizszy": "w najbliższy poniedziałek", "mianownik": "poniedziałek"},
    "Wtorek": {"na": "wtorek", "w": "we wtorek", "najblizszy": "w najbliższy wtorek", "mianownik": "wtorek"},
    "Środa": {"na": "środę", "w": "w środę", "najblizszy": "w najbliższą środę", "mianownik": "środa"},
    "Czwartek": {"na": "czwartek", "w": "w czwartek", "najblizszy": "w najbliższy czwartek", "mianownik": "czwartek"},
    "Piątek": {"na": "piątek", "w": "w piątek", "najblizszy": "w najbliższy piątek", "mianownik": "piątek"}
}

branches = {
    "-": None,
    "Ursus 1 (U1)": "U1", "Ursus 2 (U2)": "U2", "Komorów (K)": "K",
    "Michałowice (M)": "M", "Nowa Wieś (NW)": "NW", "Pruszków (P)": "P"
}
loc_grammar = {
    "U1": "Ursusa 1", "U2": "Ursusa 2", "K": "Komorowa", 
    "M": "Michałowic", "NW": "Nowej Wsi", "P": "Pruszkowa"
}

# Słownik lokalizacyjny do e-maili (gdzie: Miejscownik)
loc_grammar_w = {
    "U1": "Ursusie 1 (U1)", "U2": "Ursusie 2 (U2)", "K": "Komorowie (K)", 
    "M": "Michałowicach (M)", "NW": "Nowej Wsi (NW)", "P": "Pruszkowie (P)"
}

# MATRYCA CZASÓW DOJAZDÓW
drive_matrix = {
    "M":  {"M": 0, "K": 12, "U1": 10, "U2": 12, "NW": 12, "P": 15},
    "K":  {"M": 12, "K": 0, "U1": 20, "U2": 18, "NW": 5,  "P": 8},
    "U1": {"M": 10, "K": 20, "U1": 0, "U2": 8,  "NW": 22, "P": 18},
    "U2": {"M": 12, "K": 18, "U1": 8, "U2": 0,  "NW": 20, "P": 15},
    "NW": {"M": 12, "K": 5,  "U1": 22, "U2": 20, "NW": 0,  "P": 8},
    "P":  {"M": 15, "K": 8,  "U1": 18, "U2": 15, "NW": 8,  "P": 0}
}

# --- OBSŁUGA PLIKU ---
default_file = "Zastępstwa 2026_27 - KUM&Co.xlsx"

st.info("💡 **Wskazówka:** Aplikacja domyślnie korzysta z zapisanego na serwerze grafiku. Jeśli wiesz, że grafik uległ niedawno zmianie, warto wgrać jego zaktualizowaną wersję poniżej.")
uploaded_file = st.file_uploader("Wgraj zaktualizowany grafik (opcjonalnie)", type=["xlsx"])

if uploaded_file is not None:
    active_file = uploaded_file
    st.success("✅ Używam wgranego przez Ciebie nowszego pliku!")
elif os.path.exists(default_file):
    active_file = default_file
    st.success(f"✅ Używam domyślnego pliku z serwera: **{default_file}**.")
else:
    active_file = None
    st.warning("⚠️ Nie znaleziono domyślnego pliku na serwerze. Proszę wgrać grafik ręcznie.")

# --- USTAWIENIA NA GŁÓWNYM EKRANIE ---
st.markdown("---")
st.write("### Główne opcje")
col_opt1, col_opt2 = st.columns(2)
with col_opt1:
    my_name = st.text_input(
        "Twoje imię i nazwisko (wyklucza Cię z wyników)", 
        placeholder="np. Justyna Tymińska", 
        key="user_name", 
        on_change=update_name
    ).strip()
    st.query_params["name"] = st.session_state.user_name
    
    st.caption("💡 **Wciśnij Enter, aby zatwierdzić.** Zapisz ten adres URL w zakładkach, a aplikacja zapamięta Twoje dane!")
    
    exclude_me = my_name.lower()
    my_name_parts = exclude_me.split()
    
with col_opt2:
    st.write("") 
    is_multiple = st.checkbox("Szukam zastępstw dla więcej niż 1 grupy", value=False)
    if is_multiple:
        num_groups = st.number_input("Ile grup?", min_value=2, max_value=5, value=2)
    else:
        num_groups = 1

st.markdown("---")

# --- KONFIGURACJA GRUP ---
groups_config = []
st.write("### Parametry Grupy 1")

if "start_0" not in st.session_state:
    st.session_state["start_0"] = time(15, 50)
    st.session_state["end_0"] = time(17, 20)

col1, col2, col3 = st.columns(3)
with col1: default_day = st.selectbox("Dzień", days, key="day_0")
with col2: default_level = st.selectbox("Poziom", levels, key="level_0")
with col3: default_branch = st.selectbox("Która filia?", list(branches.keys()), key="branch_0", on_change=update_branches)

col_t1, col_t2 = st.columns(2)
with col_t1: 
    default_start = st.time_input("Czas rozpoczęcia", key="start_0", on_change=update_end_time, args=(0,))
with col_t2: 
    default_end = st.time_input("Czas zakończenia", key="end_0")

# Ostrzeżenie na dziwne godziny nocne ma priorytet
is_night_0 = default_start.hour < 7 or default_start.hour >= 22 or default_end.hour < 7 or default_end.hour >= 22
if is_night_0:
    st.warning("🦉 Nocna zmiana? O tej porze pracują tylko sowy i wampiry. Upewnij się, czy czasem nie używasz czasu w formacie 12-godzinnym.")
elif default_end < default_start:
    st.warning("🕰️ Wehikuł czasu, to byłby cud. Zajęcia kończą się przed ich rozpoczęciem. Niestety DeLorean jest w warsztacie – popraw godziny.")

groups_config.append({
    "day": default_day, "level": default_level, "branch": default_branch, 
    "start": default_start, "end": default_end
})

if is_multiple:
    for i in range(1, int(num_groups)):
        if f"start_{i}" not in st.session_state:
            st.session_state[f"start_{i}"] = st.session_state["start_0"]
            st.session_state[f"end_{i}"] = st.session_state["end_0"]
        if f"branch_{i}" not in st.session_state:
            st.session_state[f"branch_{i}"] = st.session_state.get("branch_0", list(branches.keys())[0])

        st.markdown("---")
        st.write(f"### Parametry Grupy {i+1}")
        c1, c2, c3 = st.columns(3)
        with c1: g_day = st.selectbox("Dzień", days, index=days.index(default_day), key=f"day_{i}")
        with c2: g_level = st.selectbox("Poziom", levels, key=f"level_{i}")
        with c3: g_branch = st.selectbox("Która filia?", list(branches.keys()), key=f"branch_{i}")
        
        ct1, ct2 = st.columns(2)
        with ct1: 
            g_start = st.time_input("Czas rozpoczęcia", key=f"start_{i}", on_change=update_end_time, args=(i,))
        with ct2: 
            g_end = st.time_input("Czas zakończenia", key=f"end_{i}")
            
        is_night_i = g_start.hour < 7 or g_start.hour >= 22 or g_end.hour < 7 or g_end.hour >= 22
        if is_night_i:
            st.warning("🦉 Nocna zmiana? O tej porze pracują tylko sowy i wampiry. Upewnij się, czy czasem nie używasz czasu w formacie 12-godzinnym.")
        elif g_end < g_start:
            st.warning("🕰️️ Ktoś tu chyba wynalazł wehikuł czasu! Zajęcia kończą się przed ich rozpoczęciem. Niestety DeLorean jest w warsztacie – popraw godziny.")
        
        groups_config.append({
            "day": g_day, "level": g_level, "branch": g_branch, 
            "start": g_start, "end": g_end
        })

# --- FUNKCJE POMOCNICZE CZĘŚĆ 2 ---
def overlaps(cell_time, req_start, req_end):
    return req_start < cell_time < req_end

def get_commute_warning(teacher_row, day_sheet, req_start, target_branch_code, teacher_branch_str):
    if not target_branch_code: return ""
    
    t_branches_raw = teacher_branch_str.upper().replace(" ", "").replace("&", "+").split("+")
    t_branches_valid = [b for b in t_branches_raw if b in loc_grammar]
    
    if not t_branches_valid or target_branch_code not in loc_grammar: return ""
    
    if len(t_branches_valid) == 1 and t_branches_valid[0] == target_branch_code: 
        return ""
    
    last_busy_end = None
    dummy = datetime(2000, 1, 1)
    req_dt = datetime.combine(dummy, req_start)
    
    for cell in day_sheet[4]:
        if isinstance(cell.value, time):
            col_dt = datetime.combine(dummy, cell.value)
            if col_dt <= req_dt:
                t_cell = day_sheet.cell(row=teacher_row, column=cell.column)
                if not is_cell_free(t_cell):
                    last_busy_end = col_dt 
                    
    if last_busy_end and last_busy_end <= req_dt:
        diff_mins = int((req_dt - last_busy_end).total_seconds() / 60)
        
        worst_drive_time = 0
        for b in t_branches_valid:
            d_time = drive_matrix.get(b, {}).get(target_branch_code, 0)
            if d_time > worst_drive_time:
                worst_drive_time = d_time
        
        needed_time = 15 + worst_drive_time
        
        if 0 <= diff_mins < needed_time:
            od = " / ".join([loc_grammar.get(b, b) for b in t_branches_valid])
            return f"Kończę o {last_busy_end.strftime('%H:%M')} ({diff_mins} min z: {od})"
                
    return ""

st.markdown("---")

# --- GŁÓWNY PROCES ---
if st.button("Znajdź Zastępstwo 🚀", use_container_width=True):
    if active_file is None:
        st.error("⚠ Proszę najpierw wgrać plik z grafikiem!")
    else:
        try:
            wb = openpyxl.load_workbook(active_file, data_only=True)
            if not isinstance(active_file, str): 
                active_file.seek(0)
                
            df_levels = pd.read_excel(active_file, sheet_name="Lektorzy i poziomy grup na zast", header=None)
            
            all_group_results = []
            
            for g_idx, g_conf in enumerate(groups_config):
                level_col = None
                
                for col_idx in df_levels.columns:
                    if str(df_levels.iloc[4][col_idx]).strip() == str(g_conf["level"]).strip():
                        level_col = col_idx
                        break
                        
                q_teachers = []
                if level_col is not None:
                    for idx, row in df_levels.iterrows():
                        if idx < 4: continue
                        val = str(row[level_col]).strip().lower()
                        if val == "v":
                            t_name = str(row[1]).strip()
                            
                            if exclude_me and all(part in t_name.lower() for part in my_name_parts):
                                continue 
                                
                            b_val = str(row[2]).strip()
                            phone_val = str(row[3]).strip()
                            if phone_val == "nan": phone_val = "-"
                            email = str(row[4]).strip()
                            
                            if t_name != "nan":
                                q_teachers.append({"name": t_name, "email": email, "phone": phone_val, "branch": b_val})
                                
                excel_day = day_mapping[g_conf["day"]]
                day_sheet = wb[excel_day]
                
                overlapping_cols = {}
                for cell in day_sheet[4]:
                    if isinstance(cell.value, time):
                        if overlaps(cell.value, g_conf["start"], g_conf["end"]):
                            overlapping_cols[cell.column] = cell.value
                            
                a_teachers = []
                if overlapping_cols:
                    for t in q_teachers:
                        t_row = None
                        t_name_parts = t["name"].lower().split()
                        
                        for row in range(4, day_sheet.max_row + 1):
                            c_val = str(day_sheet.cell(row=row, column=1).value).strip()
                            if c_val != "None" and all(part in c_val.lower() for part in t_name_parts):
                                t_row = row
                                break
                                
                        if t_row:
                            is_free = True
                            for col in overlapping_cols.keys():
                                if not is_cell_free(day_sheet.cell(row=t_row, column=col)):
                                    is_free = False
                                    break
                                    
                            if is_free:
                                target_b_code = branches[g_conf["branch"]]
                                commute_warn_msg = get_commute_warning(t_row, day_sheet, g_conf["start"], target_b_code, t["branch"])
                                
                                is_at_branch = False
                                if target_b_code:
                                    b_upper = t["branch"].upper()
                                    if "KUM" in b_upper and "CO" in b_upper:
                                        is_at_branch = True
                                    else:
                                        t_branches = [b.strip() for b in b_upper.replace(" ", "").replace("&", "+").split("+")]
                                        if target_b_code in t_branches:
                                            is_at_branch = True
                                            
                                a_teachers.append({
                                    "Nauczyciel": t["name"],
                                    "E-mail": t["email"],
                                    "Telefon": t["phone"],
                                    "Filie": t["branch"],
                                    "_is_at_branch": is_at_branch,
                                    "_commute_warn": commute_warn_msg
                                })
                                
                all_group_results.append(a_teachers)
                    
            st.session_state.search_results = all_group_results
            
        except Exception as e:
            st.error(f"Wystąpił błąd: {e}")

# --- WYSWIETLANIE I FILTROWANIE WYNIKÓW ---
if st.session_state.search_results is not None:
    st.write("### ⚙️ Filtry wyników")
    col_f1, col_f2, col_f3, col_f4 = st.columns(4)
    with col_f1:
        show_only_all = st.toggle("Tylko lektorzy na WSZYSTKIE grupy") if is_multiple else False
    with col_f2:
        exclude_bad_commute = st.toggle("Wyklucz lektorów z małą ilością czasu na dojazd")
    with col_f3:
        only_same_branch = st.toggle("Tylko lektorzy uczący dziś w tej filii")
    with col_f4:
        show_phones = st.toggle("Pokaż numery telefonów")

    st.info("💡 **Legenda (najedź na ikonę):** 🔥 - Wszystkie grupy | ⭐ - Część grup | 🏫 - Ta sama filia | ⚠️ - Problem z dojazdem")

    # Krok 1: Aplikujemy filtry UI na wszystkie grupy ZANIM policzymy odznaki 🔥
    filtered_results = []
    for res in st.session_state.search_results:
        g_filtered = []
        for t in res:
            if exclude_bad_commute and t["_commute_warn"] != "": continue
            if only_same_branch and not t["_is_at_branch"]: continue
            g_filtered.append(t)
        filtered_results.append(g_filtered)
        
    # Krok 2: Liczymy wystąpienia lektorów tylko w przefiltrowanej puli
    filtered_counts = {}
    for res in filtered_results:
        for t in res:
            filtered_counts[t["Nauczyciel"]] = filtered_counts.get(t["Nauczyciel"], 0) + 1

    # Krok 3: Wyświetlamy ostateczne wyniki używając poprawionych zliczeń
    for g_idx, (g_conf, res) in enumerate(zip(groups_config, filtered_results)):
        
        is_first = (g_idx == 0)
        with st.expander(f"Wyniki dla Grupy {g_idx+1} ({g_conf['day']} {g_conf['start'].strftime('%H:%M')}) - Poziom {g_conf['level']}", expanded=is_first):
            
            display_data = []
            for t in res:
                count = filtered_counts[t["Nauczyciel"]]
                
                # Ten filtr aplikujemy na samym końcu (ukrywa lektorów z małą liczbą przypisanych grup)
                if show_only_all and count < int(num_groups): continue
                    
                notes_html = []
                if t["_is_at_branch"]:
                    notes_html.append('<span title="Uczę w tej filii :)">🏫</span>')
                if t["_commute_warn"]:
                    notes_html.append(f'<span title="{t["_commute_warn"]}">⚠️</span>')
                    
                sort_all = False
                if is_multiple:
                    odmiana = "zastępstwa" if count in [2, 3, 4] else "zastępstw"
                    if count == int(num_groups):
                        notes_html.append(f'<span title="Mogę wziąć WSZYSTKIE {count} {odmiana}!">🔥</span>')
                        sort_all = True
                    elif count > 1:
                        notes_html.append(f'<span title="Mogę wziąć {count} {odmiana}.">⭐</span>')
                        
                notes_str = " ".join(notes_html) if notes_html else "-"
                
                display_data.append({
                    "Nauczyciel": t["Nauczyciel"],
                    "E-mail": t["E-mail"],
                    "Telefon": t["Telefon"],
                    "Notatki": notes_str,
                    "_sort_all": sort_all,
                    "_sort_branch": t["_is_at_branch"],
                    "_sort_warn": t["_commute_warn"] == "",
                    "_sort_count": count
                })
                
            if not display_data:
                st.warning("Brak lektorów spełniających wybrane kryteria i filtry.")
                st.warning("⚠️ **Uwaga:** Zanim wyślesz maila kryzysowego, sprawdź czy to nie przez filtry na górze, a później ręcznie w pliku Excel, czy na pewno nikogo nie ma. Ten program trzyma się na ślinę, dwa patyki i słowo honoru. Mógł się pomylić.")
                
                t_day_raw = g_conf['day']
                g_day = grammar_days[t_day_raw]
                branch_code = branches.get(g_conf['branch'])
                if branch_code:
                    t_branch_w = loc_grammar_w.get(branch_code, g_conf['branch'])
                else:
                    t_branch_w = "wybranej filii"
                t_start = g_conf['start'].strftime('%H:%M')
                t_end = g_conf['end'].strftime('%H:%M')
                
                st.markdown("#### 🚨 Generuj maila kryzysowego")
                crisis_body = f"Nikt nie może wziąć zastępstwa na {g_day['na']} ({t_start}-{t_end}) w {t_branch_w} dla grupy <TU WPISZ NAZWĘ GRUPY>.\nProszę o interwencję i wsparcie w znalezieniu zastępstwa lub odwołaniu zajęć."
                st.text_area("Skopiuj tekst poniżej:", value=crisis_body, height=120, key=f"crisis_txt_{g_idx}")
                
            else:
                display_data.sort(key=lambda x: (
                    not x["_sort_all"], 
                    not x["_sort_branch"], 
                    not x["_sort_warn"], 
                    -x["_sort_count"], 
                    x["Nauczyciel"]
                ))

                emails = [d["E-mail"] for d in display_data if d["E-mail"] != "nan" and "@" in d["E-mail"]]
                
                st.info("💡 **Pamiętaj, żeby wpisać się w tabelkę i załączyć w DW lidera, biuro i metodyków swojej filii :)**")
                
                if st.button("Przypomnij adresy e-mail", key=f"btn_emails_{g_idx}"):
                    b_name = g_conf['branch'] if g_conf['branch'] else "wybranej filii"
                    st.code(f"*email lidera w filii {b_name}*, *email biura filii {b_name}*, *e-mail metodyków filii {b_name}*", language="text")
                
                st.code("; ".join(emails), language="text")
                
                if show_phones:
                    md_table = "| Nauczyciel | E-mail | Telefon | Notatki |\n|---|---|---|---|\n"
                    for d in display_data:
                        md_table += f"| {d['Nauczyciel']} | {d['E-mail']} | {d['Telefon']} | {d['Notatki']} |\n"
                else:
                    md_table = "| Nauczyciel | E-mail | Notatki |\n|---|---|---|\n"
                    for d in display_data:
                        md_table += f"| {d['Nauczyciel']} | {d['E-mail']} | {d['Notatki']} |\n"
                    
                st.markdown(md_table, unsafe_allow_html=True)
                
                # --- GENERATOR E-MAILI ---
                st.markdown("#### ✉️ Szybka wiadomość do grupy")
                if f"tpl_{g_idx}" not in st.session_state:
                    st.session_state[f"tpl_{g_idx}"] = 0
                    
                t_day_raw = g_conf['day']
                g_day = grammar_days[t_day_raw]
                
                branch_code = branches.get(g_conf['branch'])
                if branch_code:
                    t_branch_w = loc_grammar_w.get(branch_code, g_conf['branch'])
                    t_branch_mianownik = g_conf['branch']
                else:
                    t_branch_w = "wybranej filii"
                    t_branch_mianownik = "wybranej filii"

                t_start = g_conf['start'].strftime('%H:%M')
                t_end = g_conf['end'].strftime('%H:%M')
                
                templates = [
                    f"Cześć!\nSzukam zastępstwa na {g_day['na']} ({t_start}-{t_end}) w {t_branch_w} dla grupy <TU WPISZ NAZWĘ GRUPY>.\nKtoś poratuje?\nDzięki!",
                    f"Hej wszystkim,\npotrzebuję pomocy z zastępstwem {g_day['najblizszy']}.\nZajęcia: {t_start}-{t_end} w {t_branch_w} (grupa <TU WPISZ NAZWĘ GRUPY>).\nZ góry wielkie dzięki za pomoc!",
                    f"Ratunku! Szukam dobrej duszy na zastępstwo.\nKiedy: {g_day['mianownik']}, {t_start}-{t_end}\nGdzie: {t_branch_mianownik}\nGrupa: <TU WPISZ NAZWĘ GRUPY>\nZ góry dzięki za pomoc!",
                    f"Cześć, ma ktoś może wolne okienko {g_day['w']}?\nSzukam zastępstwa w {t_branch_w} na godziny {t_start}-{t_end} (grupa <TU WPISZ NAZWĘ GRUPY>).\nOdwdzięczę się przy najbliższej okazji! :)",
                    f"Hej! Poszukiwane zastępstwo na {g_day['na']} w {t_branch_w}.\nGodziny: {t_start}-{t_end}\nGrupa: <TU WPISZ NAZWĘ GRUPY>\nPomoże ktoś?",
                    f"Cześć! Szukam zastępstwa na {g_day['na']} w {t_branch_w}. Lekcja trwa od {t_start} do {t_end} dla grupy <TU WPISZ NAZWĘ GRUPY>.\nKto da radę wziąć?\nZ góry dzięki!"
                ]
                
                current_tpl_idx = st.session_state[f"tpl_{g_idx}"]
                current_body = templates[current_tpl_idx]
                
                st.button("🎲 Losuj inny tekst", key=f"btn_rand_{g_idx}", on_click=next_tpl, args=(g_idx,))
                
                st.write("**Gotowy szablon:**")
                st.code(current_body, language="text")

# --- FOOTER Z LINKAMI I DIAGNOSTYKĄ ---
st.markdown("---")

with st.expander("dla sekretariatów: pokaż potencjalne błędy w tabelach"):
    if st.button("Uruchom diagnostykę grafiku"):
        if active_file is None:
            st.error("Wgraj najpierw plik Excel z grafikiem.")
        else:
            try:
                df_diag = pd.read_excel(active_file, sheet_name="Lektorzy i poziomy grup na zast", header=None)
                errors = []
                
                for idx, row in df_diag.iterrows():
                    if idx <= 4: continue 
                    
                    t_name = str(row[1]).strip()
                    if t_name == "nan" or not t_name: continue
                    
                    # Logika skracania odwrócona dla formatu Nazwisko Imię
                    parts = t_name.split()
                    if len(parts) > 1:
                        imiona = " ".join(parts[1:])
                        pierwsza_litera_nazwiska = parts[0][0]
                        abbr_name = f"{imiona} {pierwsza_litera_nazwiska}."
                    else:
                        abbr_name = t_name
                        
                    phone = str(row[3]).strip()
                    email = str(row[4]).strip()
                    
                    if phone == "nan" or phone == "-" or phone == "":
                        errors.append(f"**{abbr_name}** nie ma wpisanego numeru telefonu.")
                    if email == "nan" or email == "-" or email == "":
                        errors.append(f"**{abbr_name}** nie ma wpisanego maila.")
                        
                    has_levels = False
                    for col_idx in range(5, len(row)):
                        if str(row[col_idx]).strip().lower() == "v":
                            has_levels = True
                            break
                            
                    if not has_levels:
                        errors.append(f"**{abbr_name}** nie ma wpisanych żadnych poziomów jako odpowiednich.")
                        
                if errors:
                    for e in errors:
                        st.write(f"- {e}")
                else:
                    st.success("Wszystko wygląda poprawnie! Brak braków w mailach, telefonach i przypisanych poziomach.")
                    
            except Exception as e:
                st.error(f"Wystąpił błąd podczas skanowania pliku: {e}")

col_qr, col_fb = st.columns(2)
with col_qr:
    st.write("**Aplikacja na telefon (zeskanuj kod QR):**")
    st.image("https://api.qrserver.com/v1/create-qr-code/?size=150x150&data=https://tinyurl.com/needasubpls", width=120)
    st.write("Adres: [https://tinyurl.com/needasubpls](https://tinyurl.com/needasubpls)")
    
with col_fb:
    st.write("**błąd? sugestia?**")
    st.write("Wpisz tutaj: [https://tinyurl.com/cosniebangla](https://tinyurl.com/cosniebangla)")
